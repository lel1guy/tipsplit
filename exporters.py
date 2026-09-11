"""Payday outputs: payslips + cash sheet (HTML print), and the Excel/CSV escapes.

Numbers come from db + splitting — this module only formats. Vales were already
handed out during the week, so the cash that leaves the till on payday is Σ net.

print pages: /print/payslips/{week} · /print/cashsheet/{week}
data:        /api/export/week/{week}?fmt=xlsx|csv · /api/export/annual/{year}?fmt=…
"""
import csv
import io
from html import escape

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HEAD_FILL = PatternFill("solid", fgColor="1F2430")
HEAD_FONT = Font(color="FFFFFF", bold=True)

HEAD = ["Pessoa", "Função", "Horas", "Bruto", "Adiantamentos", "A pagar"]


def payable(week: dict) -> list[dict]:
    """Staff who get a line: worked hours or took an advance."""
    return [e for e in week["entries"] if e["hours"] or e["vales"]]


def week_table(week: dict) -> list[list]:
    """Header + one row per payable staff + TOTAL. Gross/pool always balance."""
    rows = [list(HEAD)]
    for e in payable(week):
        sid = e["staff_id"]
        rows.append([e["name"], e.get("position") or "", e["hours"],
                     week["gross_shares"].get(sid, 0.0), e["vales"],
                     week["shares"].get(sid, 0.0)])
    rows.append(["TOTAL", "", week["total_hours"],
                 round(sum(week["gross_shares"].values()), 2),
                 round(sum(e["vales"] for e in week["entries"]), 2),
                 round(sum(week["shares"].values()), 2)])
    return rows


def week_rows_dicts(week: dict) -> list[dict]:
    return [dict(zip(HEAD, r)) for r in week_table(week)[1:-1]]      # no TOTAL


def week_cash_total(week: dict) -> float:
    return round(sum(week["shares"].values()), 2)


# ---------- data exports ----------

def week_csv(week: dict) -> str:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")            # ; = what PT Excel expects
    w.writerows(week_table(week))
    return buf.getvalue()


def week_xlsx(week: dict) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Semana"
    _sheet(ws, week_table(week))
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out


def annual_table(report: dict) -> list[list]:
    months = report["months"]
    head = ["Pessoa", "Função"] + [f"M{m:02d}" for m in months] + ["Bruto", "Adiant.", "Total"]
    rows = [head]
    for s in sorted(report["staff"], key=lambda x: x["name"]):
        rows.append([s["name"], s["position"]]
                    + [s["monthly"].get(m, 0.0) for m in months]
                    + [s["gross"], s["vales"], s["net"]])
    return rows


def annual_csv(report: dict) -> str:
    buf = io.StringIO()
    csv.writer(buf, delimiter=";").writerows(annual_table(report))
    return buf.getvalue()


def annual_xlsx(report: dict) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = f"Gorjetas {report['year']}"
    _sheet(ws, annual_table(report))
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out


def _sheet(ws, rows: list[list]) -> None:
    for r in rows:
        ws.append(r)
    for col in range(1, len(rows[0]) + 1):
        c = ws.cell(row=1, column=col)
        c.fill, c.font = HEAD_FILL, HEAD_FONT
    ws.freeze_panes = "A2"
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        ws.column_dimensions[letter].width = min(
            max(len(str(c.value or "")) for c in col) + 2, 40)


# ---------- print pages ----------

_CSS = """
  :root { --ink:#12161a; --muted:#5b6672; --line:#c9d1d9; }
  * { box-sizing: border-box; }
  body { margin:0; padding:18px; color:var(--ink); background:#fff;
         font-family: system-ui, -apple-system, "Segoe UI", sans-serif; font-size:14px; }
  h1 { font-size:20px; margin:0 0 2px; }
  .sub { color:var(--muted); font-size:13px; margin-bottom:14px; }
  table { width:100%; border-collapse:collapse; margin-top:8px; }
  th { text-align:left; font-size:12px; color:var(--muted); border-bottom:1px solid var(--line); padding:4px 6px; }
  td { padding:4px 6px; border-bottom:1px solid #eef1f4; }
  td.num, th.num { text-align:right; font-variant-numeric:tabular-nums; }
  .total td { font-weight:700; border-top:2px solid var(--ink); }
  .statement { margin-top:10px; padding:6px 10px; border-left:3px solid #67707a;
               background:#f6f8fa; color:var(--muted); font-size:12px; }
  .sign { margin-top:26px; display:flex; justify-content:space-between; gap:40px; font-size:12px; color:var(--muted); }
  .sign span { border-top:1px solid var(--ink); padding-top:4px; min-width:190px; }
  .slip { page-break-after: always; }
  .slip:last-child { page-break-after: auto; }
  .bar { margin-bottom:14px; }
  .bar button { font-size:14px; padding:6px 12px; }
  @media print { .bar { display:none; } body { padding:0; } }
"""


def _page(title: str, body: str) -> str:
    return f"""<!DOCTYPE html>
<html lang="pt"><head><meta charset="UTF-8">
<title>{escape(title)}</title><style>{_CSS}</style></head>
<body>
<div class="bar"><button onclick="window.print()">Imprimir</button></div>
{body}
</body></html>"""


def _euros(v) -> str:
    return f"{float(v):,.2f}".replace(",", "\u00a0").replace(".", ",").replace("\u00a0", ".") + " €"


L = {
    "pt": {
        "payslip": "comprovativo de gorjetas", "week_of": "Semana de",
        "hours_worked": "Horas trabalhadas", "share": "Parte",
        "advances": "Adiantamentos (vales)", "to_receive": "A receber",
        "all_together": "com todos", "rule_full":
            "Regra: horas ÷ total de horas. Nenhuma parte fica com a casa.",
        "rule": "Regra: horas ÷ total de horas.", "signature": "Assinatura",
        "date": "Data", "cashsheet": "folha de caixa", "people": "pessoas",
        "paid_advances": "vales já pagos", "person": "Pessoa", "gross": "Bruto",
        "adv_short": "Adiant.", "to_pay": "A pagar", "paid": "Pago",
        "cash_out": "Dinheiro a tirar da caixa",
        "nothing": "Sem horas nem adiantamentos nesta semana.",
        "your_share": "a vossa parte", "average": "média", "tips": "Gorjetas",
    },
    "en": {
        "payslip": "tips payslip", "week_of": "Week of",
        "hours_worked": "Hours worked", "share": "Share",
        "advances": "Advances (vales)", "to_receive": "To receive",
        "all_together": "all together", "rule_full":
            "Rule: hours ÷ total hours. No part stays with the house.",
        "rule": "Rule: hours ÷ total hours.", "signature": "Signature",
        "date": "Date", "cashsheet": "cash sheet", "people": "people",
        "paid_advances": "advances already paid", "person": "Person",
        "gross": "Gross", "adv_short": "Adv.", "to_pay": "To pay",
        "paid": "Paid", "cash_out": "Cash out of the till",
        "nothing": "No hours or advances this week.",
        "your_share": "your share", "average": "average", "tips": "Tips",
    },
}


def _L(lang: str) -> dict:
    """Portuguese unless the venue asked for English."""
    return L["en"] if (lang or "pt").lower().startswith("en") else L["pt"]


def _slip(week: dict, e: dict, venue: str, lang: str = "pt") -> str:
    sid = e["staff_id"]
    gross = week["gross_shares"].get(sid, 0.0)
    vale = e["vales"]
    net = week["shares"].get(sid, 0.0)
    rate = week["rate_per_hour"]
    T = _L(lang)
    return f"""<div class="slip">
  <h1>{escape(venue or T["tips"])} — {T["payslip"]}</h1>
  <div class="sub">{T["week_of"]} {week["start_date"]} · {escape(e["name"])}
    {("· " + escape(e["position"])) if e.get("position") else ""}</div>
  <table>
    <tr><th>{T["hours_worked"]}</th><td class="num">{e["hours"]:g} h</td></tr>
    <tr><th>{T["share"]} ({e["hours"]:g} h ÷ {week["total_hours"]:g} h)</th>
      <td class="num">{_euros(gross)}</td></tr>
    <tr><th>{T["advances"]}</th><td class="num">− {_euros(vale)}</td></tr>
    <tr class="total"><td>{T["to_receive"]}</td><td class="num">{_euros(net)}</td></tr>
  </table>
  <div class="statement">{_euros(week["pool_eur"])} × {e["hours"]:g} h ÷
    {week["total_hours"]:g} h = {_euros(gross)}
    ({T["all_together"]}: {_euros(week["pool_eur"])} ÷ {week["total_hours"]:g} h = {_euros(rate)}/h).
    {T["rule_full"]}</div>
  <div class="sign"><span>{T["signature"]}</span><span>{T["date"]} ___/___/______</span></div>
</div>"""


def payslips_html(week: dict, venue: str = "", lang: str = "pt") -> str:
    T = _L(lang)
    slips = "".join(_slip(week, e, venue, lang) for e in payable(week))
    if not slips:
        slips = f"<p>{T['nothing']}</p>"
    title = f"{T['tips']} {week['start_date']}"
    return _page(title, slips)


def cashsheet_html(week: dict, venue: str = "", lang: str = "pt") -> str:
    T = _L(lang)
    rows = "".join(
        f"<tr><td>{escape(e['name'])}</td>"
        f"<td class='num'>{_euros(week['gross_shares'].get(e['staff_id'], 0.0))}</td>"
        f"<td class='num'>{_euros(e['vales'])}</td>"
        f"<td class='num'>{_euros(week['shares'].get(e['staff_id'], 0.0))}</td>"
        f"<td class='num'>☐</td></tr>"
        for e in payable(week))
    total = week_cash_total(week)
    body = f"""<h1>{escape(venue or T["tips"])} — {T["cashsheet"]}</h1>
<div class="sub">{T["week_of"]} {week["start_date"]} · pool {_euros(week["pool_eur"])} ·
  {len(payable(week))} {T["people"]} · {T["paid_advances"]} {_euros(sum(e["vales"] for e in week["entries"]))}</div>
<table>
  <tr><th>{T["person"]}</th><th class="num">{T["gross"]}</th><th class="num">{T["adv_short"]}</th>
      <th class="num">{T["to_pay"]}</th><th class="num">{T["paid"]}</th></tr>
  {rows}
  <tr class="total"><td>{T["cash_out"]}</td><td class="num"></td><td class="num"></td>
      <td class="num">{_euros(total)}</td><td></td></tr>
</table>
<div class="statement">{_euros(week["pool_eur"])} × horas ÷
  {week["total_hours"]:g} h = {T["your_share"]} ({T["average"]} {_euros(week["rate_per_hour"])}/h).
  {T["rule"]}</div>
<div class="sign"><span>Assinatura (gerência)</span><span>Data ___/___/______</span></div>"""
    return _page(f"Caixa {week['start_date']}", body)
