"""T-C: payday paperwork — payslips, cash sheet, week + annual exports.

The invariant that matters: the cash leaving the till is Σ net, i.e.
pool − Σ min(vale, gross). Everything else here is formatting.

Run:  pytest tests/ -q
"""
import csv
import io
import re

import pytest
from openpyxl import load_workbook

import db
import exporters
import splitting


@pytest.fixture()
def fresh(tmp_path, monkeypatch):
    monkeypatch.setattr(db, "DB_PATH", tmp_path / "t.db")
    db.init_db()
    conn = db._conn()
    for t in ("vales", "entries", "staff", "weeks", "week_pools"):
        conn.execute(f"DELETE FROM {t}")
    conn.commit()
    conn.close()
    return db


def _week(d, pool=600.0, vale_a=0.0, start="2026-10-05"):
    w = d.create_week(start)
    a = d.create_staff("Ana Test", "Bartender")
    b = d.create_staff("Bruno Test", "Barback")
    d.save_week(w["id"], pool, [
        {"staff_id": a["id"], "mon": 8, "tue": 8, "wed": 8, "thu": 8, "fri": 8},
        {"staff_id": b["id"], "mon": 4, "tue": 4, "wed": 4, "thu": 4, "fri": 4},
    ])
    if vale_a:
        d.record_vale(a["id"], vale_a, week_id=w["id"])
    return d.get_week(w["id"]), a, b


def test_cash_total_is_pool_minus_the_advances(fresh):
    w, _, _ = _week(fresh, 600.0, vale_a=25.0)
    # Ana 400−25 + Bruno 200 = 575 = pool − vales
    assert exporters.week_cash_total(w) == pytest.approx(575.0, abs=0.01)


def test_cash_total_with_advance_above_gross(fresh):
    """A vale over someone's gross can't come out of the till twice."""
    w, _, _ = _week(fresh, 600.0, vale_a=500.0)     # Ana gross 400
    capped = sum(min(e["vales"], w["gross_shares"][e["staff_id"]]) for e in w["entries"])
    assert exporters.week_cash_total(w) == pytest.approx(round(600 - capped, 2), abs=0.01)
    assert exporters.week_cash_total(w) == 200.0    # only Bruno is handed cash


def test_week_table_balances(fresh):
    w, a, b = _week(fresh, 555.0, vale_a=20.0)
    rows = exporters.week_table(w)
    assert rows[0] == exporters.HEAD
    assert len(rows) == 4                            # header + 2 staff + TOTAL
    assert rows[1][0] == "Ana Test"
    assert rows[1][3] == pytest.approx(370.0, abs=0.01)   # 555 × 40/60
    assert rows[1][5] == pytest.approx(350.0, abs=0.01)   # 370 − 20 vale
    total = rows[-1]
    assert total[0] == "TOTAL" and total[3] == 555.0  # gross balances to the pool
    assert total[5] == pytest.approx(535.0, abs=0.01)


def test_zero_hours_zero_vale_staff_are_skipped(fresh):
    w = fresh.create_week("2026-10-12")
    ghost = fresh.create_staff("Ghost Test")
    fresh.save_week(w["id"], 100.0, [{"staff_id": ghost["id"], "mon": 0}])
    w2 = fresh.get_week(w["id"])
    assert exporters.payable(w2) == []
    assert "Ghost" not in exporters.week_csv(w2)


def test_week_csv_parses(fresh):
    w, _, _ = _week(fresh, 300.0)
    rows = list(csv.reader(io.StringIO(exporters.week_csv(w)), delimiter=";"))
    assert rows[0] == exporters.HEAD and len(rows) == 4
    assert float(rows[1][2]) == 40.0                  # hours
    assert float(rows[1][3]) == 200.0                 # gross of the 300 pool


def test_week_xlsx_opens_with_expected_values(fresh):
    w, _, _ = _week(fresh, 600.0, vale_a=10.0)
    ws = load_workbook(exporters.week_xlsx(w)).active
    assert [c.value for c in ws[1]] == exporters.HEAD
    assert {ws.cell(row=r, column=1).value for r in (2, 3)} == {"Ana Test", "Bruno Test"}
    last = ws.max_row
    assert ws.cell(row=last, column=1).value == "TOTAL"
    assert float(ws.cell(row=last, column=4).value) == 600.0
    assert float(ws.cell(row=last, column=6).value) == pytest.approx(590.0, abs=0.01)


def test_annual_buckets_by_month_and_totals(fresh):
    w1, a, _ = _week(fresh, 600.0, start="2026-10-05")
    w2 = fresh.create_week("2026-11-02")
    fresh.save_week(w2["id"], 400.0, [{"staff_id": a["id"], "mon": 8, "tue": 8}])
    rep = fresh.annual(2026)
    assert rep["weeks"] == 2 and rep["months"] == [10, 11]
    row = [s for s in rep["staff"] if s["name"] == "Ana Test"][0]
    assert row["monthly"][10] == pytest.approx(400.0, abs=0.01)
    assert row["monthly"][11] == pytest.approx(400.0, abs=0.01)
    assert row["net"] == pytest.approx(800.0, abs=0.01)
    assert row["position"] == "Bartender"


def test_annual_ignores_other_years(fresh):
    _week(fresh)                                       # 2026-10-05
    assert fresh.annual(2025) == {"year": 2025, "weeks": 0, "months": [], "staff": []}


def test_annual_exports(fresh):
    _week(fresh, 100.0)
    rep = fresh.annual(2026)
    assert "Ana Test" in exporters.annual_csv(rep)
    ws = load_workbook(exporters.annual_xlsx(rep)).active
    assert ws.cell(row=1, column=1).value == "Pessoa"
    assert "M10" in [c.value for c in ws[1]]


def test_payslips_and_cashsheet_carry_the_numbers(fresh):
    w, a, b = _week(fresh, 600.0, vale_a=25.0)
    slips = exporters.payslips_html(w, "Bar Teste")
    assert "Bar Teste" in slips and "Ana Test" in slips and "Assinatura" in slips
    assert "375,00" in slips                          # Ana's net, PT formatting
    sheet = exporters.cashsheet_html(w, "Bar Teste")
    assert "Dinheiro a tirar da caixa" in sheet and "Ana Test" in sheet
    assert sheet.count("☐") == 2                      # one tick box per payable person
    assert "575,00" in sheet                          # total leaving the till


def test_payslip_statement_arithmetic_reconciles(fresh):
    """The printed proof must survive a calculator: pool × hours ÷ total = share."""
    w, a, b = _week(fresh, 466.0)                      # Ana 40h, Bruno 20h → 60h
    html = exporters.payslips_html(w, "Bar Teste")
    m = re.search(r"([\d.]+,\d\d) € × ([\d.,]+) h ÷\s*([\d.,]+) h = ([\d.]+,\d\d) €", html)
    assert m, "statement is not in the exact 'pool × hours ÷ total = share' form"

    def num(s):
        return float(s.replace(".", "").replace(",", ".")) if "," in s else float(s)

    pool, hours, total, shown = (num(m.group(i)) for i in (1, 2, 3, 4))
    assert total == w["total_hours"] == 60.0
    assert abs(round(pool * hours / total, 2) - shown) < 0.01, (pool, hours, total, shown)
    assert abs(shown - w["gross_shares"][a["id"]]) < 0.01


def test_week_statement_is_a_rate_not_a_multiplication(fresh):
    """The week line states €/h only — nothing on paper rounds a rate then multiplies it."""
    w, _, _ = _week(fresh, 466.0)
    st = splitting.statement(w["pool_eur"], w["total_hours"])
    assert "÷" in st and "/h" in st and "Regra" in st
    assert "×" not in st


def test_settings_roundtrip(fresh):
    assert fresh.get_setting("venue_name", "") == ""
    fresh.set_setting("venue_name", "Bar Teste")
    assert fresh.get_setting("venue_name") == "Bar Teste"
    fresh.set_setting("venue_name", "Bar Dois")
    assert fresh.get_setting("venue_name") == "Bar Dois"


def test_printed_pages_follow_the_venue_language(fresh):
    """settings.lang drives the printouts. PT stays byte-identical; EN is the mirror."""
    full, _, _ = _week(fresh, pool=600.0, vale_a=20.0)
    pt = exporters.payslips_html(full, "Bar Teste")
    en = exporters.payslips_html(full, "Bar Teste", "en")
    for pt_label in ("Horas trabalhadas", "A receber", "Assinatura", "Regra: horas ÷ total"):
        assert pt_label in pt, pt_label
        assert pt_label not in en, pt_label
    for en_label in ("Hours worked", "To receive", "Signature", "Rule: hours ÷ total hours"):
        assert en_label in en, en_label
    assert "€" in en and "," in en          # euro format stays PT: 33,60 €

    cs_pt = exporters.cashsheet_html(full, "Bar Teste")
    cs_en = exporters.cashsheet_html(full, "Bar Teste", "en")
    assert "Dinheiro a tirar da caixa" in cs_pt and "folha de caixa" in cs_pt
    assert "Cash out of the till" in cs_en and "cash sheet" in cs_en
    assert "Dinheiro a tirar da caixa" not in cs_en

    # the spreadsheet exports stay Portuguese on purpose: they are data files for the
    # accountant and their row mapping is keyed on HEAD


def test_user_facing_errors_follow_the_venue_language(fresh, monkeypatch):
    import main
    fresh.set_setting("lang", "en")
    assert main._msg("semana fechada", "week locked") == "week locked"
    fresh.set_setting("lang", "pt")
    assert main._msg("semana fechada", "week locked") == "semana fechada"
